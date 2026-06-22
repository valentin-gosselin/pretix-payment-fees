"""
Tests for the "Détail par commande" exporter and its renderers (STORY-201).

Checks registration, the 3 formats, the GDPR guarantee (no email in any output)
and FR labels under a forced locale.
"""
import io
from decimal import Decimal

import pytest
from django.utils import translation
from django.utils.timezone import now
from django_scopes import scope, scopes_disabled

from pretix_payment_fees.exporters.order_detail import OrderDetailExporter
from pretix_payment_fees.renderers.order_detail_tabular import (
    OrderDetailCSVRenderer,
    OrderDetailExcelRenderer,
    column_headers,
)
from pretix_payment_fees.services.order_detail_builder import OrderDetailBuilder


@pytest.fixture
def event(db):
    from pretix.base.models import (
        Event, Item, Order, OrderFee, OrderPosition, Organizer,
    )

    with scopes_disabled():
        org = Organizer.objects.create(name="Org", slug="ode-org")
        ev = Event.objects.create(
            organizer=org, name="ODE Event", slug="ev-ode",
            date_from=now(), currency="EUR", plugins="pretix_payment_fees",
        )
        item = Item.objects.create(
            event=ev, name="Place", default_price=Decimal("20.00")
        )
        web = org.sales_channels.get(identifier="web")
        o = Order.objects.create(
            code="ZZZZZ", event=ev, status=Order.STATUS_PAID,
            datetime=now(), expires=now(), total=Decimal("20.00"),
            sales_channel=web, email="secret@example.com",
        )
        OrderPosition.objects.create(order=o, item=item, price=Decimal("20"))
        OrderFee.objects.create(
            order=o, fee_type=OrderFee.FEE_TYPE_PAYMENT,
            internal_type="mollie_creditcard_fee", value=Decimal("0.50"),
        )
    return ev


def _exp(event):
    return OrderDetailExporter(event, event.organizer)


def test_registered(event):
    from pretix.base.signals import register_data_exporters

    with scope(organizer=event.organizer):
        ids = [getattr(r, "identifier", None)
               for _r, r in register_data_exporters.send(event) if r]
    assert "order_detail_psp" in ids


def test_render_pdf(event):
    with scope(organizer=event.organizer):
        fn, ct, data = _exp(event).render({"_format": "pdf"})
    assert fn.endswith(".pdf") and ct == "application/pdf"
    assert data[:4] == b"%PDF"


def test_render_csv_bom(event):
    with scope(organizer=event.organizer):
        fn, ct, data = _exp(event).render({"_format": "csv"})
    assert fn.endswith(".csv") and data[:3] == b"\xef\xbb\xbf"


def test_render_xlsx(event):
    with scope(organizer=event.organizer):
        fn, ct, data = _exp(event).render({"_format": "xlsx"})
    assert fn.endswith(".xlsx") and data[:2] == b"PK"


def test_no_email_in_any_output(event):
    """GDPR: the buyer email must not leak into PDF/CSV/Excel."""
    with scope(organizer=event.organizer):
        _f, _c, pdf = _exp(event).render({"_format": "pdf"})
        _f, _c, csv = _exp(event).render({"_format": "csv"})
        _f, _c, xls = _exp(event).render({"_format": "xlsx"})
    assert b"secret@example.com" not in pdf
    assert b"secret@example.com" not in csv
    assert b"secret@example.com" not in xls


def test_csv_french_headers(event):
    with scope(organizer=event.organizer), translation.override("fr"):
        report = OrderDetailBuilder([event]).build()
        head = column_headers(report)
    assert "Commande" in head
    assert "Produit" in head
    assert "Montant" in head


def test_excel_numeric_total(event):
    import openpyxl

    with scope(organizer=event.organizer):
        _f, _c, data = _exp(event).render({"_format": "xlsx"})
    ws = openpyxl.load_workbook(io.BytesIO(data)).active
    # grand total row: gross column (6) should be a number
    gross = ws.cell(row=ws.max_row, column=6).value
    assert float(gross) == 20.0
