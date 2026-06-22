"""
CSV/Excel renderers for the "Remplissage de salle" export (STORY-202).

Two sections (categories, quotas) in one sheet/file. No monetary value.
"""
import csv
import io

from django.utils.translation import gettext_lazy as _

from openpyxl import Workbook
from openpyxl.styles import Border, Font, PatternFill, Side

BOM = "﻿"
ACCENT = "4F46E5"
HEAD_FILL = PatternFill("solid", fgColor="EEF0FB")
THIN = Side(style="thin", color="E5E7EB")


def _rows(report):
    """Build labelled rows for both sections (plain strings/ints, no money)."""
    cat_head = [str(_("Catégorie")), str(_("Payant")),
                str(_("Invitations")), str(_("Total"))]
    rows = [("section", str(_("Ventes par catégorie"))), ("head", cat_head)]
    for c in report.categories:
        rows.append(("row", [c.name, c.paid, c.free, c.total]))
    rows.append(("total", [str(_("Total")), report.total_paid,
                           report.total_free, report.total_places]))
    rows.append(("blank", []))
    show_rate = report.has_quota_sizes
    if show_rate:
        q_head = [str(_("Quota")), str(_("Vendu")),
                  str(_("Capacité")), str(_("Remplissage"))]
    else:
        q_head = [str(_("Quota")), str(_("Vendu"))]
    rows.append(("section", str(_("Taux de remplissage"))))
    rows.append(("head", q_head))
    for q in report.quotas:
        if show_rate:
            cap = q.size if q.size else ""
            rows.append(("row", [q.name, q.sold, cap, q.rate_display]))
        else:
            rows.append(("row", [q.name, q.sold]))
    return rows


class RemplissageCSVRenderer:
    def __init__(self, report):
        self.report = report

    def render(self) -> bytes:
        out = io.StringIO()
        out.write(BOM)
        w = csv.writer(out, delimiter=";", lineterminator="\r\n")
        for kind, payload in _rows(self.report):
            if kind == "blank":
                w.writerow([])
            elif kind == "section":
                w.writerow([str(payload)])
            else:
                w.writerow([str(x) for x in payload])
        return out.getvalue().encode("utf-8")


class RemplissageExcelRenderer:
    def __init__(self, report):
        self.report = report

    def render(self) -> bytes:
        wb = Workbook()
        ws = wb.active
        ws.title = str(_("Remplissage"))
        for kind, payload in _rows(self.report):
            if kind == "blank":
                ws.append([])
                continue
            if kind == "section":
                ws.append([str(payload)])
                ws.cell(row=ws.max_row, column=1).font = Font(
                    bold=True, color=ACCENT, size=11)
                continue
            ws.append([x for x in payload])
            r = ws.max_row
            if kind == "head":
                for c in range(1, len(payload) + 1):
                    cell = ws.cell(row=r, column=c)
                    cell.font = Font(bold=True, color=ACCENT, size=9)
                    cell.fill = HEAD_FILL
                    cell.border = Border(
                        bottom=Side(style="medium", color=ACCENT))
            elif kind == "total":
                for c in range(1, len(payload) + 1):
                    ws.cell(row=r, column=c).font = Font(bold=True, color=ACCENT)
            else:
                for c in range(1, len(payload) + 1):
                    ws.cell(row=r, column=c).border = Border(bottom=THIN)
        for col in "ABCD":
            ws.column_dimensions[col].width = 24
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()
