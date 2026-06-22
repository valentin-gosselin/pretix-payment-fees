"""
Excel renderer for the "Recette Manifestation" export (STORY-105).

Serialises the same flattened structure as the CSV renderer (recette_tabular),
so totals are identical across CSV, Excel and the PDF. Amounts are written as
real numbers with a EUR cell format, headers and total rows are emphasised.
"""
from decimal import Decimal

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from .recette_tabular import (
    KIND_GRAND,
    KIND_SUBTOTAL,
    KIND_TOTAL,
    column_headers,
    flatten,
)

EUR_FMT = "# ##0.00 €"
ACCENT = "4F46E5"
HEADER_FILL = PatternFill("solid", fgColor="EEF0FB")
TOTAL_FILL = PatternFill("solid", fgColor="EEF0FB")
THIN = Side(style="thin", color="E5E7EB")


class RecetteExcelRenderer:
    """Render a RecetteReport to XLSX bytes."""

    def __init__(self, report):
        self.report = report

    def render(self) -> bytes:
        wb = Workbook()
        ws = wb.active
        ws.title = "Recette"

        headers = column_headers(self.report)
        ncols = len(headers)
        # index of the trailing "kind" column (not displayed as data)
        kind_idx = ncols - 1
        # numeric columns: unit price (6), gross (7), fee cols.., net (-2)
        amount_cols = self._amount_columns(ncols)

        ws.append(headers[:kind_idx])  # drop the technical "kind" column
        self._style_header(ws, kind_idx)

        for row in flatten(self.report):
            kind = row[-1]
            values = list(row[:kind_idx])
            ws.append(values)
            r = ws.max_row
            self._style_row(ws, r, kind, amount_cols, kind_idx)

        self._autosize(ws, kind_idx)
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    # -- helpers ------------------------------------------------------------

    @staticmethod
    def _amount_columns(ncols):
        """1-based column indices holding amounts (unit price, gross, fees, net).

        Layout: channel, session, product, nature, vat, count, unit_price,
        gross, <fees...>, net, kind. Amounts = unit_price(7), gross(8),
        fees..(9..n-2), net(n-1).
        """
        # count fixed leading cols: 6 (channel..count) then unit_price at 7
        return list(range(7, ncols))  # 7..(ncols-1) covers unit_price..net

    def _style_header(self, ws, kind_idx):
        for c in range(1, kind_idx + 1):
            cell = ws.cell(row=1, column=c)
            cell.font = Font(bold=True, color=ACCENT, size=9)
            cell.fill = HEADER_FILL
            cell.border = Border(bottom=Side(style="medium", color=ACCENT))
            cell.alignment = Alignment(
                horizontal="right" if c >= 6 else "left", vertical="center"
            )
        ws.freeze_panes = "A2"

    def _style_row(self, ws, r, kind, amount_cols, kind_idx):
        is_total = kind in (KIND_TOTAL, KIND_GRAND)
        is_sub = kind == KIND_SUBTOTAL
        for c in range(1, kind_idx + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = Border(bottom=THIN)
            if c in amount_cols and isinstance(cell.value, (int, float, Decimal)):
                cell.number_format = EUR_FMT
                cell.alignment = Alignment(horizontal="right")
            elif c == 6:  # count
                cell.alignment = Alignment(horizontal="right")
            if is_total:
                cell.font = Font(bold=True, color=ACCENT)
                cell.fill = TOTAL_FILL
            elif is_sub:
                cell.font = Font(bold=True)

    def _autosize(self, ws, kind_idx):
        for c in range(1, kind_idx + 1):
            width = 10
            for row in ws.iter_rows(min_col=c, max_col=c):
                v = row[0].value
                if v is not None:
                    width = max(width, min(len(str(v)) + 2, 40))
            ws.column_dimensions[
                ws.cell(row=1, column=c).column_letter
            ].width = width
